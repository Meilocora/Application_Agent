"""Word documents: cover letters and one report per search run - plain code, no LLM involved.

File names start with the date (YYYYMMDD_...), so the output folder sorts chronologically.
"""

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph

from application_agent.config import OUTPUT_DIR
from application_agent.models import CandidateProfile, CoverLetter, ScoredJob, SearchPreferences

_RED = RGBColor(0xB0, 0x1E, 0x1E)
_GREY = RGBColor(0x59, 0x59, 0x59)


def _output_path(name: str) -> Path:
    """``OUTPUT_DIR/YYYYMMDD_<name>.docx`` - with _2, _3, ... instead of overwriting an existing file."""
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    # Keep only letters, digits, - and _ so the file name is valid on every OS.
    stem = f"{date.today():%Y%m%d}_{re.sub(r'[^\w-]+', '_', name)[:80].strip('_')}"
    path = OUTPUT_DIR / f"{stem}.docx"
    counter = 2
    while path.exists():
        path = OUTPUT_DIR / f"{stem}_{counter}.docx"
        counter += 1
    return path


def _new_document():
    document = Document()
    section = document.sections[0]
    # DIN 5008 margins - harmless for the report as well.
    section.left_margin, section.right_margin = Cm(2.5), Cm(2)
    section.top_margin, section.bottom_margin = Cm(2), Cm(2)
    normal = document.styles["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(11)
    normal.paragraph_format.space_after = Pt(0)
    return document


def _add_hyperlink(paragraph: Paragraph, text: str, url: str) -> None:
    """Append a clickable link. python-docx has no API for this, so the XML is built by hand."""
    relation_id = paragraph.part.relate_to(url, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relation_id)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend([color, underline])
    run.append(properties)
    text_element = OxmlElement("w:t")
    text_element.text = text
    run.append(text_element)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


# ---------------------------------------------------------------------------
# Cover letter
# ---------------------------------------------------------------------------


def save_cover_letter(
    letter: CoverLetter, profile: CandidateProfile, preferences: SearchPreferences, scored: ScoredJob
) -> Path:
    """Save the letter as .docx. The sender block is built here, so the address never reaches the LLM."""
    document = _new_document()

    sender = document.add_paragraph()
    sender.add_run(profile.name).bold = True
    # "Street 1, 12345 City" -> one line per part
    for line in (part.strip() for part in preferences.address.split(",")):
        sender.add_run(f"\n{line}")
    sender.paragraph_format.space_after = Pt(36)

    recipient = document.add_paragraph("\n".join(filter(None, [scored.job.company, scored.job.location])))
    recipient.paragraph_format.space_after = Pt(24)

    date_line = document.add_paragraph(f"{date.today():%d.%m.%Y}")
    date_line.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    date_line.paragraph_format.space_after = Pt(24)

    subject = document.add_paragraph()
    subject.add_run(letter.subject).bold = True
    subject.paragraph_format.space_after = Pt(18)

    for text in [letter.salutation, *letter.paragraphs, letter.closing]:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.space_after = Pt(10)
    document.add_paragraph().paragraph_format.space_after = Pt(18)  # room for the signature
    document.add_paragraph(profile.name)

    path = _output_path(f"cover_letter_{scored.job.company}_{scored.job.title}")
    document.save(path)
    return path


# ---------------------------------------------------------------------------
# Search report
# ---------------------------------------------------------------------------


def _add_labeled(document, label: str, text: str, color: RGBColor | None = None) -> None:
    paragraph = document.add_paragraph()
    paragraph.add_run(f"{label}: ").bold = True
    run = paragraph.add_run(text)
    if color:
        run.font.color.rgb = color
    paragraph.paragraph_format.space_after = Pt(4)


def _add_job(document, number: int, scored: ScoredJob, is_new: bool, preferences: SearchPreferences) -> None:
    job, match = scored.job, scored.match
    document.add_heading(f"{number}. {job.title}", level=2)

    facts = [f"{job.company}", job.location, f"Score {match.score} ({match.verdict})"]
    fastest = scored.fastest_commute
    if fastest:
        over = " - over your limit" if fastest[1] > preferences.max_commute_minutes else ""
        facts.append(f"{fastest[1]} min by {fastest[0]}{over}")
    meta = document.add_paragraph(" | ".join(filter(None, facts)))
    meta.paragraph_format.space_after = Pt(2)

    id_line = document.add_paragraph()
    id_run = id_line.add_run(f"ID: {job.id}" + ("" if is_new else "   (already found in an earlier run)"))
    id_run.font.size, id_run.font.color.rgb = Pt(9), _GREY
    id_line.paragraph_format.space_after = Pt(6)

    if not match.fits_employment_type:
        _add_labeled(document, "Warning", "not the employment type you are looking for", _RED)
    summary = document.add_paragraph(match.summary)
    summary.paragraph_format.space_after = Pt(6)
    if match.matching_skills:
        _add_labeled(document, "Matching skills", ", ".join(match.matching_skills))
    if match.missing_skills:
        _add_labeled(document, "Missing skills", ", ".join(match.missing_skills))
    if match.selling_points:
        _add_labeled(document, "Highlight in your letter", "")
        for point in match.selling_points:
            document.add_paragraph(point, style="List Bullet")
    if match.red_flags:
        _add_labeled(document, "Red flags", "; ".join(match.red_flags), _RED)

    links = document.add_paragraph()
    links.paragraph_format.space_before = Pt(6)
    if job.application_url:
        _add_hyperlink(links, "Apply", job.application_url)
        links.add_run("   |   ")
    _add_hyperlink(links, "Job ad", job.url)


def save_search_report(
    scored_jobs: list[ScoredJob], new_ids: set[str], strategy: str, preferences: SearchPreferences
) -> Path:
    """One report per search run with all postings of this run, best first."""
    document = _new_document()
    document.add_heading(f"Job search {date.today():%d.%m.%Y}", level=1)

    intro = document.add_paragraph(
        f"{len(scored_jobs)} matching postings, {len(new_ids)} of them new. "
        f"Looking for: {preferences.employment_type} | {', '.join(preferences.target_roles)}"
    )
    intro.paragraph_format.space_after = Pt(6)
    _add_labeled(document, "Search strategy", strategy)
    hint = document.add_paragraph("Write a cover letter with: uv run cover-letter <ID>")
    hint.runs[0].font.color.rgb = _GREY

    for number, scored in enumerate(scored_jobs, start=1):
        _add_job(document, number, scored, scored.job.id in new_ids, preferences)

    path = _output_path("job_search")
    document.save(path)
    return path
