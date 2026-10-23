"""Excel view of the stored postings. Rebuilt from the database after every run - edits in the file are lost."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from application_agent.config import EXCEL_PATH
from application_agent.models import JobStatus
from application_agent.storage import StoredJob, load_jobs

# (header, column width) - the order here is the column order in the sheet.
_COLUMNS = [
    ("ID", 22),
    ("Status", 11),
    ("Score", 8),
    ("Verdict", 10),
    ("Title", 40),
    ("Company", 28),
    ("Location", 22),
    ("Commute (min)", 10),
    ("Commute details", 26),
    ("Summary", 70),
    ("Red flags", 45),
    ("Apply", 12),
    ("Job ad", 12),
    ("Found", 18),
]
_HEADERS = [header for header, _ in _COLUMNS]
_WRAPPED = {"Title", "Summary", "Red flags", "Commute details"}
_HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
_HEADER_FONT = Font(bold=True, color="FFFFFF")


def _row(stored: StoredJob) -> list:
    scored = stored.scored
    job, match = scored.job, scored.match
    fastest = scored.fastest_commute
    return [
        job.id,
        stored.status,
        match.score,
        match.verdict,
        job.title,
        job.company,
        job.location,
        fastest[1] if fastest else None,
        ", ".join(f"{mode}: {minutes}" for mode, minutes in scored.commute_minutes.items()),
        match.summary,
        "; ".join(match.red_flags),
        "Apply" if job.application_url else "",
        "Job ad" if job.url else "",
        stored.found_at.replace("T", " "),
    ]


def export_excel(path: Path = EXCEL_PATH) -> Path | None:
    """Write all stored postings to ``path``. Returns None if the file is open in Excel."""
    # Open postings on top, then by score - load_jobs already sorts by score.
    jobs = sorted(load_jobs(), key=lambda stored: stored.status != JobStatus.NEW)

    workbook = Workbook()
    sheet = workbook.active
    # Tells the type checker what a new workbook always has: one normal worksheet.
    # (`active` is typed "Worksheet, Chartsheet or None".)
    assert isinstance(sheet, Worksheet)
    sheet.title = "Jobs"
    sheet.append(_HEADERS)
    for cell in sheet[1]:
        cell.fill, cell.font = _HEADER_FILL, _HEADER_FONT

    # Determine the column indices for the "Apply" and "Job ad" columns.
    apply_col, ad_col = _HEADERS.index("Apply") + 1, _HEADERS.index("Job ad") + 1
    for stored in jobs:
        sheet.append(_row(stored))
        row = sheet.max_row
        # Add hyperlinks for the "Apply" and "Job ad" columns if URLs are available.
        for column, url in ((apply_col, stored.scored.job.application_url), (ad_col, stored.scored.job.url)):
            if url:
                cell = sheet.cell(row=row, column=column)
                assert isinstance(cell, Cell)  # typed "Cell or MergedCell"; we never merge cells
                cell.hyperlink, cell.style = url, "Hyperlink"

    # Adjust column widths and wrap text for specified columns.
    for index, (header, width) in enumerate(_COLUMNS, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        if header in _WRAPPED:
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        else:
            for cell in sheet[letter][1:]:
                cell.alignment = Alignment(vertical="top")

    # Freeze the top row so headers are always visible.
    sheet.freeze_panes = "A2"
     # Apply autofilter to the entire sheet.
    sheet.auto_filter.ref = sheet.dimensions
    if jobs:  # red -> yellow -> green for the score
        score_col = get_column_letter(_HEADERS.index("Score") + 1)
        sheet.conditional_formatting.add(
            f"{score_col}2:{score_col}{sheet.max_row}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                           mid_type="num", mid_value=50, mid_color="FFEB84",
                           end_type="num", end_value=100, end_color="63BE7B"),
        )

    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(path)
    except PermissionError:  # Windows locks files that are open in Excel
        return None # prevents crash if the file is open in Excel
    return path
